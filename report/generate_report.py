from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / 'results'
FIG = ROOT / 'figures'
OUT = ROOT / 'report'
OUT.mkdir(exist_ok=True)

summary = pd.read_csv(RES/'main_results_summary.csv').set_index('algorithm')
stats = pd.read_csv(RES/'wilcoxon_holm_balanced_accuracy_with_r.csv')
fried = pd.read_csv(RES/'friedman_balanced_accuracy.csv').iloc[0]
thr = pd.read_csv(RES/'deployment_thresholds_top2.csv').set_index('algorithm')
ab = pd.read_csv(RES/'scaling_ablation_summary.csv').set_index('algorithm')
final_cfg = pd.read_csv(RES/'final_configurations.csv').set_index('algorithm')
audit = json.load(open(RES/'data_audit.json'))
env = json.load(open(RES/'environment.json'))

DOCX = OUT / 'Abah_Emmanuel_Enemaku_CPE513_P1_Report.docx'

def pct(x): return f"{100*x:.2f}%"
def pm_pct(m,s): return f"{100*m:.2f} ± {100*s:.2f}%"
def pm(x,s,dec=3): return f"{x:.{dec}f} ± {s:.{dec}f}"
def pformat(p):
    if p < .001: return f"{p:.2e}"
    return f"{p:.3f}"

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)

styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(11)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
styles['Normal'].paragraph_format.line_spacing=1.15
styles['Normal'].paragraph_format.space_after=Pt(6)
for sname,size,color in [('Title',20,'17365D'),('Heading 1',14,'17365D'),('Heading 2',11.5,'17365D')]:
    st=styles[sname]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(4); st.paragraph_format.keep_with_next=True
if 'CaptionCustom' not in [s.name for s in styles]:
    cap=styles.add_style('CaptionCustom',WD_STYLE_TYPE.PARAGRAPH)
    cap.font.name='Arial'; cap.font.size=Pt(9); cap.font.italic=True
    cap._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); cap._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    cap.paragraph_format.space_before=Pt(2); cap.paragraph_format.space_after=Pt(6)
if 'TableText' not in [s.name for s in styles]:
    tt=styles.add_style('TableText',WD_STYLE_TYPE.PARAGRAPH)
    tt.font.name='Arial'; tt.font.size=Pt(8.3)
    tt._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); tt._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    tt.paragraph_format.space_after=Pt(0); tt.paragraph_format.line_spacing=1.0

# footer page numbers
for section in doc.sections:
    p=section.footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Page ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)


def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False, align=None, fontsize=8.3):
    cell.text=''
    p=cell.paragraphs[0]; p.style=styles['TableText']
    if align is not None: p.alignment=align
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(fontsize); r.font.name='Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table_caption(text):
    p=doc.add_paragraph(style='CaptionCustom'); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); r.bold=True; r.italic=False
    return p

def add_table(headers, rows, widths=None, fontsize=8.3):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i],h,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,fontsize=fontsize)
        set_cell_shading(t.rows[0].cells[i],'D9EAF7')
        if widths: t.rows[0].cells[i].width=Inches(widths[i])
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i],val,align=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER,fontsize=fontsize)
            if widths: cells[i].width=Inches(widths[i])
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    return t

def add_fig(path, caption, width=5.8):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width))
    c=doc.add_paragraph(caption,style='CaptionCustom'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

def body(text, bold_prefix=None):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def heading(text, level=1):
    return doc.add_heading(text,level=level)

# ---------------- title page ----------------
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(60)
r=p.add_run('EMPIRICAL COMPARISON OF LEARNING ALGORITHMS\nFOR BREAST CANCER DIAGNOSIS'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(23,54,93); r.font.name='Arial'
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10)
r=p.add_run('Problem P1 - Breast Cancer Wisconsin (Diagnostic)'); r.bold=True; r.font.size=Pt(13)

details=[
 ('Course','CPE 513 - Artificial Neural Network'),
 ('Student','Abah Emmanuel Enemaku'),
 ('Student number','2021/1/79515CP'),
 ('Department','Computer Engineering'),
 ('Institution','Federal University of Technology, Minna'),
 ('Dataset','UCI Breast Cancer Wisconsin (Diagnostic), dataset 17; file: wdbc.data'),
 ('Dataset DOI','10.24432/C5DW2B'),
 ('Date','18 September 2026'),
 ('Code archive / GitHub','Submission archive supplied with report; GitHub profile: https://github.com/Enemaku'),
]
t=doc.add_table(rows=len(details),cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,(a,b) in enumerate(details):
    set_cell_text(t.rows[i].cells[0],a,bold=True,fontsize=9.5); set_cell_text(t.rows[i].cells[1],b,fontsize=9.5)
    t.rows[i].cells[0].width=Inches(1.55); t.rows[i].cells[1].width=Inches(4.6)
    for c in t.rows[i].cells:
        tcPr=c._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
        if borders is None:
            borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
        for e in ('top','left','bottom','right','insideH','insideV'):
            edge=OxmlElement(f'w:{e}'); edge.set(qn('w:val'),'nil'); borders.append(edge)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12)
r=p.add_run('Abstract'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(23,54,93)
abs_text=(
"Breast mass classification from fine-needle aspirate measurements is a small-sample diagnostic support problem in which false-negative errors are clinically important. This study asks whether performance differences among four algorithm families are larger than fold-to-fold variation under an identical evaluation protocol. The UCI Breast Cancer Wisconsin (Diagnostic) dataset (569 cases, 30 numeric predictors) was evaluated with Logistic Regression, k-nearest neighbours, Random Forest and an RBF Support Vector Machine (SVM) using repeated stratified 5-fold cross-validation (three repeats) with nested 3-fold tuning and equal six-configuration search budgets. RBF SVM achieved 97.38 ± 1.42% balanced accuracy, 99.58 ± 0.54% ROC-AUC and 98.28 ± 2.57% sensitivity at 95% specificity. Logistic Regression was statistically indistinguishable in balanced accuracy (97.19 ± 1.43%; Holm-adjusted paired p = 1.00), while both exceeded Random Forest and k-NN in the paired comparison. Removing fold-wise scaling reduced balanced accuracy by 6.56 percentage points for SVM and 4.04 points for k-NN. A provisional SVM threshold of 0.291 is preferred for screening support, subject to external clinical validation."
)
p=doc.add_paragraph(abs_text); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(4)
p=doc.add_paragraph(); p.add_run('Keywords: ').bold=True; p.add_run('breast cancer diagnosis; repeated cross-validation; support vector machine; logistic regression; random forest; k-nearest neighbours; ROC-AUC; model comparison')
doc.add_page_break()

# ---------------- 1 Introduction ----------------
heading('1 Introduction')
body('The decision supported by this model is whether a breast mass sampled by fine-needle aspiration should be classified as malignant or benign from quantitative characteristics of cell nuclei. The output is best viewed as diagnostic decision support for a clinician rather than an autonomous clinical decision. In this study malignant disease is the positive class, so the main practical concern is preserving sensitivity while constraining false-positive referrals.')
body('Formally, the input is x in R^30 and the target y is binary (malignant = 1, benign = 0). The comparison question is: under identical outer folds, preprocessing, tuning budget and metric definitions, are differences in balanced accuracy among Logistic Regression, k-NN, Random Forest and RBF SVM larger than fold-to-fold variation, and what sensitivity is available when specificity is constrained to approximately 95%?')
body('The anchor study by Aamir et al. [2] compared Random Forest, Gradient Boosting, SVM, an artificial neural network and a multilayer perceptron on WDBC after correlation-based feature elimination and recursive feature elimination reduced the representation to 11 attributes. Their best reported 80:20 result was 99.12% accuracy for the MLP; Random Forest and SVM were reported at 98.07% and 97.76%, respectively. The study also describes a 5-fold framework, but the feature-selection procedure, use of several split ratios and emphasis on accuracy make the headline numbers different in protocol from the present experiment.')
body('The contribution here is not a new classifier. It is a protocol-controlled comparison using the same 30-feature representation for every algorithm, preprocessing fitted only within folds, an equal six-configuration tuning budget per model, repeated paired outer folds, multiple clinically relevant metrics, a paired statistical comparison, timing, calibration and an explicit scaling ablation.')

# ---------------- 2 Related work ----------------
heading('2 Related work')
body('Recent WDBC studies typically report high scores, but their conclusions are difficult to compare directly because feature selection, tuning effort, train/test splits and headline metrics vary. Aamir et al. [2] combine feature selection with model comparison. Rasool et al. [3] also combine exploratory feature elimination and hyperparameter optimization, reporting particularly strong polynomial-SVM results. Dhahri et al. [4] automate preprocessing, feature selection and classifier optimization using genetic programming. These studies establish that the dataset is highly learnable, while also motivating a comparison in which the algorithm family is the main controlled variable.')
add_table_caption('Table 1. Prior work on the Wisconsin diagnostic breast-cancer classification problem.')
rows=[
 ['Aamir et al. [2]','WDBC; reduced to 11 features','RF, GB, SVM, ANN, MLP','5-fold framework; 60:40, 70:30 and 80:20 results','Not fully specified per model','MLP 99.12% accuracy; RF 98.07%; SVM 97.76% (80:20 column)'],
 ['Rasool et al. [3]','WDBC; feature analysis/elimination','Polynomial/linear SVM, LR, k-NN, voting','Train/test evaluation plus k-fold CV','Hyperparameters optimized; equal budget not stated','Polynomial SVM 98.25% test accuracy; CV headline 99.03%; LR 98.06%'],
 ['Dhahri et al. [4]','Wisconsin diagnostic data','LR, k-NN, RF, GB, SVM, trees, Bayes, LDA, ensembles','10-fold CV; 5-fold used for ROC plots','Genetic-programming search across pipelines','Automated pipeline validation accuracy 98.24%; RF/GB/ET strong by log-loss'],
]
add_table(['Study','Dataset / representation','Algorithms','Validation','Tuning budget','Headline result'],rows,widths=[0.85,1.1,1.15,1.15,1.0,1.5],fontsize=7.2)
body('The gap addressed here is therefore experimental rather than algorithmic: the cited studies do not settle whether small differences among common model families remain after folds, representation, search budget and metrics are controlled and after fold-level uncertainty is made visible. The present design keeps those quantities fixed and treats the fold scores as paired observations.')

# ---------------- 3 Problem and data ----------------
heading('3 Problem and data')
body('The Breast Cancer Wisconsin (Diagnostic) dataset was obtained from the UCI Machine Learning Repository [1]. It contains measurements computed from digitised images of fine-needle aspirates. Ten nuclear characteristics (including radius, texture, perimeter, area, compactness, concavity and concave points) are each represented by mean, standard-error and worst-value summaries, producing 30 real-valued predictors. The raw file also contains an identifier and the diagnosis label; the identifier is excluded before modelling.')
add_table_caption('Table 2. Dataset summary and audit.')
rows=[
 ['Source','UCI Machine Learning Repository, dataset 17'],['File / DOI','wdbc.data / 10.24432/C5DW2B'],['Licence','CC BY 4.0'],['Size','569 cases; 30 numeric predictors after ID removal'],['Target',f"Malignant: {audit['malignant']} ({audit['malignant_percent']:.1f}%); Benign: {audit['benign']} ({100-audit['malignant_percent']:.1f}%)"],['Missing / sentinel values','0 missing values; no documented sentinel values in the 30 predictors'],['Duplicates',f"{audit['duplicate_feature_rows']} duplicate predictor rows"],['Leakage control','Identifier removed; all scaling fitted within training folds'],
]
add_table(['Item','Finding'],rows,widths=[1.45,4.85],fontsize=8.5)
body(f"The audit found no missing values and no duplicated predictor rows. Dependence among predictors was substantial: {audit['high_corr_pairs_abs_r_ge_0.90']} feature pairs had absolute Pearson correlation at least 0.90. The strongest pair was mean radius versus mean perimeter (|r| = {audit['max_abs_correlation']:.3f}); mean radius, mean perimeter and mean area, and their worst-value counterparts, formed the expected highly correlated groups. This collinearity is relevant to coefficient stability, so individual Logistic Regression coefficients are not interpreted as independent biological effects.")
body('The evaluation uses repeated stratified 5-fold cross-validation with three repeats (15 paired outer test scores per algorithm) and fixed seed 513. Each outer training partition contains its own inner 3-fold search. No model choice, hyperparameter or preprocessing statistic is selected from an outer test fold. There is no separate final test set; the outer folds are the untouched evaluation folds. The dataset is de-identified, but its age, small size and lack of demographic variables prevent any claim of subgroup fairness or prospective clinical validity.')

# ---------------- 4 Methods ----------------
heading('4 Methods')
body('Logistic Regression is included as the high-bias linear/probabilistic baseline [5]. k-NN supplies an instance-based method whose Euclidean neighbourhoods make feature scale directly visible [6]. Random Forest represents a nonlinear ensemble of decorrelated decision trees [7]. RBF SVM supplies a kernel method capable of nonlinear decision boundaries through a radial basis similarity function [8]. Implementations use scikit-learn [10].')
body('For the main experiment the pipeline is: outer split -> fit StandardScaler on the outer-training data only -> inner grid search on the scaled training folds -> refit the selected configuration on the complete outer-training partition -> evaluate once on the untouched outer test fold. The same outer fold indices are reused for all four algorithms. Scaling is applied to all four pipelines to keep the numerical representation identical; for trees this monotonic transformation does not change the ordering used by splits, while for k-NN and SVM it prevents large-magnitude features from dominating distance calculations.')
add_table_caption('Table 3. Hyperparameter tuning protocol. Each algorithm receives six candidate configurations.')
rows=[
 ['Logistic Regression','C = {0.01, 0.1, 1, 10, 100, 300}','Grid','6','3-fold stratified','ROC-AUC'],
 ['k-NN','k = {3, 7, 15}; weights = {uniform, distance}','Grid','6','3-fold stratified','ROC-AUC'],
 ['Random Forest','trees = {30, 80}; max depth = {None, 8, 16}; leaf = 1','Grid','6','3-fold stratified','ROC-AUC'],
 ['RBF SVM','C = {0.1, 1, 10}; gamma = {scale, 0.01}','Grid','6','3-fold stratified','ROC-AUC'],
]
add_table(['Algorithm','Search space','Method','Configs','Inner CV','Selection'],rows,widths=[1.15,2.25,0.55,0.55,0.85,0.8],fontsize=7.8)
body('Four headline metrics are reported. Balanced accuracy is the mean of sensitivity and specificity, giving both classes equal weight. ROC-AUC measures ranking discrimination across thresholds [11]. Sensitivity at 95% specificity is the largest test-fold sensitivity available at a ROC operating point with specificity at least 0.95; it is an evaluation metric and is not reused as a deployment threshold. Matthews correlation coefficient (MCC) summarises all four cells of the confusion matrix on a -1 to +1 scale [12]. Brier score is additionally used for calibration comparison. Every headline value is the mean ± sample standard deviation across the 15 outer scores.')
body('The inferential comparison uses a Friedman test on paired balanced-accuracy scores across the four algorithms, followed by pairwise Wilcoxon signed-rank tests with Holm correction when the omnibus null is rejected [9]. Effect size r is reported from the normal approximation as |Z|/sqrt(15). A secondary ablation repeats the complete nested evaluation for k-NN and RBF SVM with StandardScaler removed; this is deliberately separate from the main comparison.')
body(f"The reference run used Python {env['python'].split()[0]}, scikit-learn {env['scikit_learn']}, NumPy {env['numpy']}, pandas {env['pandas']}, SciPy {env['scipy']} and Matplotlib {env['matplotlib']} on a Linux cloud CPU runtime ({env.get('cpu_model','CPU not reported')}, {env.get('cpu_count','?')} logical CPUs allocated, approximately {env.get('ram_gb','?')} GB RAM). A GPU is not used by these scikit-learn estimators. Training time is the wall time for the inner search plus refit on each outer fold; prediction latency is measured on the corresponding outer test fold.")

# ---------------- 5 Results ----------------
heading('5 Results')
body('Table 4 reports the repeated outer-fold results. RBF SVM had the highest mean balanced accuracy, ROC-AUC and sensitivity at the 95%-specificity operating point. Logistic Regression was very close on all four headline metrics. Random Forest had the largest tuning-and-fit time in this implementation.')
add_table_caption('Table 4. Main results over 15 paired outer folds (mean ± SD). Best mean is bold in each column (higher is better for performance; lower is better for time).')
order=['RBF SVM','Logistic Regression','k-NN','Random Forest']
rows=[]
for a in order:
    s=summary.loc[a]
    rows.append([a,pm_pct(s.balanced_accuracy_mean,s.balanced_accuracy_std),pm_pct(s.roc_auc_mean,s.roc_auc_std),pm_pct(s.sensitivity_at_95_specificity_mean,s.sensitivity_at_95_specificity_std),pm(s.mcc_mean,s.mcc_std,3),f"{s.train_time_s_mean:.3f} ± {s.train_time_s_std:.3f}",f"{s.prediction_time_ms_per_sample_mean:.4f} ± {s.prediction_time_ms_per_sample_std:.4f}"])
t4 = add_table(['Algorithm','Balanced acc.','ROC-AUC','Sens. @ 95% spec.','MCC','Train (s)','Pred. (ms/sample)'],rows,widths=[1.1,1.05,0.9,1.15,0.8,0.85,1.1],fontsize=7.5)
# Mark the best value in every result column: maximize performance metrics, minimize cost.
for ri, ci in [(1,1),(1,2),(1,3),(1,4),(3,5),(2,6)]:
    for run in t4.cell(ri,ci).paragraphs[0].runs:
        run.bold = True
add_fig(FIG/'figure1_fold_level_dispersion.png','Figure 1. Fold-level dispersion for the four required metrics under repeated stratified 5-fold cross-validation (15 scores per algorithm). Triangles denote means.',width=5.75)
body('At the default probability threshold, the pooled outer-fold predictions in Figure 2 contain three predictions for each of the 569 cases because each case appears once as test data in each repeat. RBF SVM produced 25 false-negative and 14 false-positive predictions across the 1,707 pooled decisions; Logistic Regression produced 31 false-negative and 8 false-positive predictions.')
add_fig(FIG/'figure2_confusion_matrices_top2.png','Figure 2. Pooled outer-fold confusion matrices for the two highest-ranked algorithms at the default 0.5 threshold. Counts sum to 1,707 per algorithm because the 569 cases are tested in three repeats.',width=5.75)

body(f"The Friedman test rejected equality of balanced-accuracy ranks (chi-square = {fried['statistic']:.2f}, p = {fried['p_value']:.2e}). Pairwise results are shown in Table 5. The 0.19-percentage-point difference between RBF SVM and Logistic Regression did not survive the paired test, whereas both models exceeded Random Forest and k-NN after Holm correction.")
add_table_caption('Table 5. Paired statistical comparison of balanced accuracy. Effect size r = |Z|/sqrt(15); Holm-adjusted p-values control the six post-hoc comparisons.')
rows=[[f"Friedman (4 alg.)",f"chi2={fried['statistic']:.2f}",pformat(fried['p_value']),'--','--','Reject equal ranks']]
for _,r in stats.iterrows():
    comp=f"{r.algorithm_a} vs {r.algorithm_b}"
    delta=f"{100*r.mean_difference_a_minus_b:+.2f} pp"
    interp='Significant' if r.significant_0_05 else 'Not significant'
    rows.append([comp,f"W={r.wilcoxon_W:g}; Δ={delta}",pformat(r.p_raw),pformat(r.p_holm),f"{r.effect_size_r:.3f}",interp])
add_table(['Test / pair','Statistic / difference','Raw p','Holm p','Effect r','Interpretation'],rows,widths=[1.5,1.35,0.7,0.7,0.65,0.95],fontsize=7.3)

body('Figure 3 shows calibration from pooled outer-fold probabilities for RBF SVM and Logistic Regression. Their mean outer-fold Brier scores were 0.0191 ± 0.0084 and 0.0218 ± 0.0070, respectively.')
add_fig(FIG/'figure3_calibration_top2.png','Figure 3. Calibration curves for RBF SVM and Logistic Regression from pooled outer-fold probabilities; the dashed line denotes perfect calibration.',width=4.65)

add_table_caption('Table 6. Candidate operating thresholds for the two highest-ranked algorithms. Thresholds are derived from 5-fold out-of-fold predictions on the full development dataset after tuning and therefore require external validation before clinical use.')
rows=[]
for a in ['RBF SVM','Logistic Regression']:
    trow=thr.loc[a]
    sout=summary.loc[a]
    rows.append([a,f"{trow.proposed_deployment_threshold:.3f}",pct(trow.full_data_oof_sensitivity),pct(trow.full_data_oof_specificity),pm_pct(sout.sensitivity_at_95_specificity_mean,sout.sensitivity_at_95_specificity_std),f"{sout.brier_score_mean:.4f} ± {sout.brier_score_std:.4f}"])
add_table(['Algorithm','Candidate threshold','OOF sensitivity','OOF specificity','Outer sens. @95% spec.','Brier'],rows,widths=[1.45,1.05,1.0,1.0,1.35,0.85],fontsize=7.8)

body('The scaling ablation produced the summary in Table 7. Removing StandardScaler lowered mean balanced accuracy for both scale-sensitive methods under the same outer folds and tuning-budget size.')
add_table_caption('Table 7. Scaling ablation (balanced accuracy, mean ± SD over 15 outer folds).')
rows=[]
for a in ['k-NN','RBF SVM']:
    z=ab.loc[a]
    rows.append([a,pm_pct(z.scaled_balanced_accuracy_mean,z.scaled_balanced_accuracy_std),pm_pct(z.unscaled_balanced_accuracy_mean,z.unscaled_balanced_accuracy_std),f"{100*z.delta_scaled_minus_unscaled:+.2f} pp"])
add_table(['Algorithm','Scaled inside fold','No scaling','Difference'],rows,widths=[1.6,1.6,1.6,1.2],fontsize=8.3)

# ---------------- 6 Discussion ----------------
heading('6 Discussion')
body('The central result is that RBF SVM and Logistic Regression form a statistically indistinguishable top pair under this protocol. Their balanced accuracies differ by only 0.19 percentage points (97.38% versus 97.19%), while each model has a fold-to-fold standard deviation of about 1.4 percentage points. The paired SVM-versus-Logistic comparison gives Holm-adjusted p = 1.00 and effect size r = 0.036, so the data do not support declaring either one the accuracy winner. In contrast, the differences between either top model and Random Forest or k-NN are around 1.4-1.7 points and survive Holm correction with large paired effect sizes in this 15-score experiment.')
body('The ranking partly agrees with Aamir et al. [2] in that SVM and Random Forest were both competitive on WDBC, but their strongest model was an MLP and their reported accuracy values cannot be compared numerically with the balanced-accuracy results here. The most important protocol differences are their correlation/RFE feature selection to 11 attributes, the use of several train-test ratios, and accuracy as the main headline measure. Here all 30 predictors are retained, model search is nested inside repeated outer folds, and each algorithm receives six candidate configurations. A lower Random Forest balanced accuracy in the present experiment therefore does not contradict their 98.07% accuracy; the target quantity and experimental design differ before the numbers are compared.')
body('The geometry of the data helps explain why a simple linear model remains close to the nonlinear SVM. Many WDBC predictors are strongly correlated transformations of the same underlying size and contour characteristics, and the classes are separable enough that a regularised linear boundary is already strong. The RBF SVM can bend that boundary locally and gains a small average advantage, but the gain is smaller than the resampling variation. Random Forest is less sensitive to scaling and can model nonlinear interactions, but on only 569 observations its tree partitions also introduce sampling variance. k-NN depends directly on local distances and therefore pays a stronger penalty when irrelevant scale differences distort neighbourhoods.')
body(f"The scaling ablation makes that mechanism measurable. Mean balanced accuracy fell from {pct(ab.loc['k-NN'].scaled_balanced_accuracy_mean)} to {pct(ab.loc['k-NN'].unscaled_balanced_accuracy_mean)} for k-NN (a {100*ab.loc['k-NN'].delta_scaled_minus_unscaled:.2f}-point loss) and from {pct(ab.loc['RBF SVM'].scaled_balanced_accuracy_mean)} to {pct(ab.loc['RBF SVM'].unscaled_balanced_accuracy_mean)} for RBF SVM (a {100*ab.loc['RBF SVM'].delta_scaled_minus_unscaled:.2f}-point loss). This is why scaling is kept inside the fold rather than applied once to the full dataset: it is both influential and a fitted transformation.")
add_fig(FIG/'figure4_scaling_ablation.png','Figure 4. Secondary scaling ablation for k-NN and RBF SVM. Error bars are fold-level standard deviations across the same 15 outer splits.',width=4.9)
body(f"For a screening-support operating point, I would provisionally use the RBF SVM with a malignancy-probability threshold near {thr.loc['RBF SVM'].proposed_deployment_threshold:.2f}, rather than the default 0.50. In the full-data out-of-fold threshold analysis this candidate produced {pct(thr.loc['RBF SVM'].full_data_oof_sensitivity)} sensitivity and {pct(thr.loc['RBF SVM'].full_data_oof_specificity)} specificity. The corresponding Logistic Regression candidate threshold ({thr.loc['Logistic Regression'].proposed_deployment_threshold:.2f}) produced {pct(thr.loc['Logistic Regression'].full_data_oof_sensitivity)} sensitivity and {pct(thr.loc['Logistic Regression'].full_data_oof_specificity)} specificity. The SVM choice is based on the operating point and slightly lower Brier score, not on a claim of statistically superior balanced accuracy. The threshold is a development-set candidate only; a real deployment would require prospective external validation, explicit clinical cost analysis and recalibration where needed.")

# ---------------- 7 Threats ----------------
heading('7 Threats to validity')
body('Internal validity. Leakage is reduced by placing scaling and model selection inside the appropriate training folds and by reusing identical outer splits across algorithms. Remaining threats include the finite six-configuration grids and implementation-specific probability estimation for SVM. Repeated cross-validation also reuses observations across repeats, so the 15 scores are paired but not independent samples of different patients.')
body('External validity. WDBC is a small historical dataset from one diagnostic setting. It contains no demographic covariates and does not represent changes in imaging, sampling, clinical workflow or population mix that would occur in modern deployment. The reported scores therefore support a benchmark comparison, not a claim of clinical effectiveness.')
body('Construct validity. Balanced accuracy, ROC-AUC, MCC and sensitivity at a fixed-specificity target capture different aspects of classifier performance, but none encodes the real clinical cost of a missed malignancy versus an unnecessary referral. The chosen 95% specificity target is an assignment-defined operating constraint, not a validated clinical policy.')
body('Conclusion validity. Only one dataset is studied and N = 15 paired fold scores is modest. Holm correction is applied to the six post-hoc balanced-accuracy comparisons, and differences within the fold-level dispersion are treated as inconclusive rather than ranked. The candidate threshold is estimated from development-data out-of-fold predictions and remains optimistic until tested externally.')

# ---------------- 8 Conclusion ----------------
heading('8 Conclusion')
body('Under repeated stratified 5-fold cross-validation with nested equal-budget tuning, RBF SVM and Logistic Regression were the strongest models for WDBC, with balanced accuracy of 97.38 ± 1.42% and 97.19 ± 1.43%, respectively. Their 0.19-point difference was not reliable after paired testing, while both exceeded Random Forest and k-NN on balanced accuracy. For screening support, the RBF SVM provides the stronger operating-point result, with 98.28 ± 2.57% sensitivity at 95% specificity and a provisional threshold near 0.29. The strongest caveat is that these are internal benchmark results on 569 historical cases; external prospective validation is required before any clinical use.')

# ---------------- 9 References ----------------
heading('9 References')
refs=[
'[1] W. Wolberg, O. Mangasarian, N. Street, and W. Street, "Breast Cancer Wisconsin (Diagnostic) [Dataset]," UCI Machine Learning Repository, 1993. doi: 10.24432/C5DW2B.',
'[2] S. Aamir, A. Rahim, Z. Aamir, et al., "Predicting Breast Cancer Leveraging Supervised Machine Learning Techniques," Computational and Mathematical Methods in Medicine, vol. 2022, Art. no. 5869529, 2022. doi: 10.1155/2022/5869529.',
'[3] A. Rasool, C. Bunterngchit, L. Tiejian, M. R. Islam, Q. Qu, and Q. Jiang, "Improved Machine Learning-Based Predictive Models for Breast Cancer Diagnosis," International Journal of Environmental Research and Public Health, vol. 19, no. 6, Art. no. 3211, 2022. doi: 10.3390/ijerph19063211.',
'[4] H. Dhahri, E. Al Maghayreh, A. Mahmood, W. Elkilani, and M. F. Nagi, "Automated Breast Cancer Diagnosis Based on Machine Learning Algorithms," Journal of Healthcare Engineering, vol. 2019, Art. no. 4253641, 2019. doi: 10.1155/2019/4253641.',
'[5] D. R. Cox, "The Regression Analysis of Binary Sequences," Journal of the Royal Statistical Society: Series B, vol. 20, no. 2, pp. 215-242, 1958.',
'[6] T. Cover and P. Hart, "Nearest Neighbor Pattern Classification," IEEE Transactions on Information Theory, vol. 13, no. 1, pp. 21-27, 1967. doi: 10.1109/TIT.1967.1053964.',
'[7] L. Breiman, "Random Forests," Machine Learning, vol. 45, pp. 5-32, 2001. doi: 10.1023/A:1010933404324.',
'[8] C. Cortes and V. Vapnik, "Support-Vector Networks," Machine Learning, vol. 20, pp. 273-297, 1995. doi: 10.1007/BF00994018.',
'[9] J. Demsar, "Statistical Comparisons of Classifiers over Multiple Data Sets," Journal of Machine Learning Research, vol. 7, pp. 1-30, 2006.',
'[10] F. Pedregosa et al., "Scikit-learn: Machine Learning in Python," Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.',
'[11] T. Fawcett, "An Introduction to ROC Analysis," Pattern Recognition Letters, vol. 27, no. 8, pp. 861-874, 2006. doi: 10.1016/j.patrec.2005.10.010.',
'[12] B. W. Matthews, "Comparison of the Predicted and Observed Secondary Structure of T4 Phage Lysozyme," Biochimica et Biophysica Acta, vol. 405, no. 2, pp. 442-451, 1975. doi: 10.1016/0005-2795(75)90109-9.'
]
for ref in refs:
    p=doc.add_paragraph(ref); p.paragraph_format.left_indent=Cm(0.4); p.paragraph_format.first_line_indent=Cm(-0.4); p.paragraph_format.space_after=Pt(3); p.alignment=WD_ALIGN_PARAGRAPH.LEFT

# ---------------- appendices ----------------
heading('Appendix A - Reproducibility')
body('The complete submission package contains the experiment script, a Colab launcher notebook, pinned environment files, raw fold-level CSV files, fold-specific selected hyperparameters, out-of-fold predictions and all figures used in this report. The official UCI wdbc.data URL is requested by the script; an offline fallback uses scikit-learn\'s packaged copy of the same 569 x 30 predictor matrix.')
add_table_caption('Table A1. Reproducibility information.')
rows=[
 ['Seed','513'],['Primary command','python src/run_experiment.py'],['Environment','requirements.txt and environment.yml'],['Raw main results','results/fold_results_main.csv'],['Outer predictions','results/oof_predictions_main.csv'],['Fold hyperparameters','results/best_hyperparameters_by_fold.csv'],['Statistical tests','results/friedman_balanced_accuracy.csv; results/wilcoxon_holm_balanced_accuracy_with_r.csv'],['Ablation results','results/fold_results_no_scaling_ablation.csv'],['Figures','figures/figure1_... through figure4_...'],['Code delivery','Archive submitted with the report; GitHub profile: https://github.com/Enemaku'],
]
add_table(['Item','Value / path'],rows,widths=[1.6,4.7],fontsize=8.5)
body('A marker can regenerate the evidence from a clean environment by installing the pinned dependencies and running the single command above. The script writes raw per-fold results before summary tables are derived. Timing is hardware-dependent; predictive metrics should reproduce within deterministic numerical tolerance under the pinned versions.')

heading('Appendix B - AI-use statement')
ai=(
'I used ChatGPT (OpenAI) as an AI assistant to help interpret the assignment requirements, design a leakage-safe nested cross-validation workflow, draft and review Python code, organise the statistical analysis, and edit the wording and layout of this report. The assistant was also used to locate the official UCI dataset metadata and the cited anchor/related papers. All numerical results in the report come from executable experiment code and archived CSV outputs; they were not invented or copied from the assistant. I verified the dataset dimensions and class counts, inspected the preprocessing and fold boundaries, checked that each algorithm received the same tuning-budget size, and cross-checked the report values against the generated CSV files. The supplied notebook and script can regenerate the tables and figures. I remain responsible for understanding, explaining and reproducing the submitted work.'
)
body(ai)

heading('Appendix C - Final configurations')
body('The main reported scores use fold-specific configurations selected independently inside each outer-training partition. For a final model fitted after the comparison, a separate 5-fold search on all development data selected the configurations below. These full-data selections are provided for reproducibility; they do not replace the nested estimates in Table 4.')
add_table_caption('Table C1. Full-development-data selected configurations.')
rows=[]
for a in ['Logistic Regression','k-NN','Random Forest','RBF SVM']:
    cfg=final_cfg.loc[a]
    extra='StandardScaler -> model'
    if a=='RBF SVM': extra+='; probability=True'
    rows.append([a,cfg.full_data_best_params,f"{cfg.full_data_inner_roc_auc:.4f}",extra])
add_table(['Algorithm','Selected hyperparameters','5-fold inner ROC-AUC','Pipeline'],rows,widths=[1.35,2.65,1.0,1.3],fontsize=7.2)
body(f"For threshold-sensitive use, the candidate full-development thresholds are {thr.loc['RBF SVM'].proposed_deployment_threshold:.3f} for RBF SVM and {thr.loc['Logistic Regression'].proposed_deployment_threshold:.3f} for Logistic Regression. These values must be re-estimated or externally validated if the data distribution or prevalence changes.")

doc.save(DOCX)
print(DOCX)
